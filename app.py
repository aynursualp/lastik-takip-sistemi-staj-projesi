from flask import Flask, render_template, request, redirect, flash, session
import sqlite3
from lastik_kodlari import lastik_tipi_kodla
import math
import joblib
import pandas as pd
from datetime import datetime
from collections import Counter
import re
from flask import send_file
import io
import json
from openpyxl.utils import get_column_letter
from werkzeug.security import generate_password_hash, check_password_hash
from sklearn.model_selection import train_test_split
from sklearn.ensemble import RandomForestRegressor
from sklearn.metrics import mean_absolute_error
from email_servisi import musteriye_uyari_maili_gonder
from sms_servisi import sms_gonder


app = Flask(__name__)

app.secret_key = "gizli_anahtar"

model = joblib.load('lastik_tahmin_modeli.joblib')

def arac_durumu_hesapla(c, arac_id, model):
    tahmini_km = None
    kalan_km = None
    durum = None
    dusuk_guven = False
    sebep = None

    c.execute("SELECT islem_tarihi, kilometre, lastik_tipi FROM Islemler WHERE arac_id = ? ORDER BY islem_tarihi DESC LIMIT 1",(arac_id,))
    son_islem = c.fetchone()

    if not son_islem:
        return tahmini_km, kalan_km, durum, dusuk_guven, sebep
    
    islem_tarihi_str, kilometre, lastik_tipi_str = son_islem
    lastik_tipi_say = lastik_tipi_kodla(lastik_tipi_str)
    if lastik_tipi_say is None:
        lastik_tipi_say = 0
    
    islem_tarihi = pd.to_datetime(islem_tarihi_str)
    bugun = pd.to_datetime(datetime.today().strftime('%Y-%m-%d'))
    gun_farki = (bugun - islem_tarihi).days

    soru = pd.DataFrame({'kilometre':[kilometre], 'lastik_tipi':[lastik_tipi_say], 'gun_farki':[gun_farki]})
    tahmini_fark = model.predict(soru)[0]
    tahmini_km = int(kilometre + tahmini_fark)
    kalan_km = tahmini_km - kilometre

    c.execute("SELECT islem_tarihi FROM Islemler WHERE arac_id = ? ORDER BY islem_tarihi ASC",(arac_id,))
    tum_tarihler = [pd.to_datetime(r[0]) for r in c.fetchall()]

    arac_islem_sayisi = len(tum_tarihler)
    dusuk_guven = arac_islem_sayisi < 3

    ortalama_gun_araligi = None
    if len(tum_tarihler) >= 2:
        gun_farklari = [(tum_tarihler[i+1] - tum_tarihler[i]).days for i in range(len(tum_tarihler) - 1)]
        ortalama_gun_araligi = sum(gun_farklari) / len(gun_farklari)

    zaman_gecikmis = False 
    zaman_yaklasiyor = False
    if ortalama_gun_araligi:
        if gun_farki >= ortalama_gun_araligi:
            zaman_gecikmis = True
        elif gun_farki >= 0.7 * ortalama_gun_araligi:
            zaman_yaklasiyor = True
    
    km_kritik = kalan_km <= 1000
    km_yaklasiyor = kalan_km <= 3000

    if km_kritik or zaman_gecikmis:
        durum = 'kritik'
        sebep = 'km' if km_kritik else 'zaman'
    elif km_yaklasiyor or zaman_yaklasiyor:
        durum = 'yaklasiyor'
        sebep = 'km' if km_yaklasiyor else 'zaman'
    else:
        durum = 'normal'
    
    return tahmini_km, kalan_km, durum, dusuk_guven, sebep

@app.route('/')
def ana_sayfa():
    if 'giris_yapildi' not in session:
        return redirect('/login')
    
    gelen_arama = request.args.get('arama_kelimesi')
    if gelen_arama and len(gelen_arama.strip()) < 2:
        gelen_arama = None
    sayfa = request.args.get('page', 1, type=int)
    sayfa_basina_kayit = 6
    offset_degeri = (sayfa - 1) * sayfa_basina_kayit

    sirala = request.args.get('sirala', 'eski')
    siralama_haritasi = {
        'eski': 'Musteriler.musteri_id ASC',
        'yeni': 'Musteriler.musteri_id DESC',
        'ad_az': 'Musteriler.ad COLLATE NOCASE ASC',
        'ad_za': 'Musteriler.ad COLLATE NOCASE DESC'
    }

    siralama_sql = siralama_haritasi.get(sirala, 'Musteriler.musteri_id ASC')

    conn = sqlite3.connect('database.db')
    c = conn.cursor()

    c.execute('SELECT COUNT(*) FROM Musteriler')
    toplam_musteri = c.fetchone()[0]

    c.execute('SELECT COUNT(*) FROM Araclar')
    toplam_arac = c.fetchone()[0]

    c.execute('SELECT COUNT(*) FROM Islemler')
    toplam_islem = c.fetchone()[0]

    if gelen_arama:
        aranan_sart = f"%{gelen_arama}%"
        c.execute('''
                  SELECT COUNT(DISTINCT Musteriler.musteri_id)
                  FROM Musteriler
                  LEFT JOIN Araclar ON Musteriler.musteri_id = Araclar.musteri_id
                  WHERE Musteriler.ad LIKE ? OR Araclar.plaka LIKE ? OR Araclar.marka LIKE ?
                  ''',(aranan_sart, aranan_sart, aranan_sart))
        kayit_sayisi = c.fetchone()[0]
    else:
        kayit_sayisi = toplam_musteri

    toplam_sayfa = math.ceil(kayit_sayisi / sayfa_basina_kayit) if kayit_sayisi > 0 else 1

    if gelen_arama:
        aranan_sart = f"%{gelen_arama}%"

        c.execute(f'''
                  SELECT
                    Musteriler.ad,
                    Musteriler.soyad,
                    Musteriler.telefon,
                    Araclar.marka,
                    Araclar.plaka,
                    Musteriler.musteri_id,
                    Araclar.arac_id
                  FROM Musteriler
                  LEFT JOIN Araclar ON Musteriler.musteri_id = Araclar.musteri_id
                  WHERE Musteriler.ad LIKE ? OR Araclar.plaka LIKE ? OR Araclar.marka LIKE ?
                  ORDER BY {siralama_sql}
                  LIMIT ? OFFSET ?''', (aranan_sart, aranan_sart, aranan_sart, sayfa_basina_kayit, offset_degeri)
                  )
    else:
        c.execute(f'''
                  SELECT 
                    Musteriler.ad,
                    Musteriler.soyad,
                    Musteriler.telefon,
                    Araclar.marka,
                    Araclar.plaka,
                    Musteriler.musteri_id,
                    Araclar.arac_id
                  FROM Musteriler
                  LEFT JOIN Araclar ON Musteriler.musteri_id = Araclar.musteri_id
                  ORDER BY {siralama_sql}
                  LIMIT ? OFFSET ?''', (sayfa_basina_kayit, offset_degeri)
                  )

    if gelen_arama:
        aranan_sart = f"%{gelen_arama}%"
        c.execute(f'''
                  SELECT DISTINCT Musteriler.musteri_id
                  FROM Musteriler
                  LEFT JOIN Araclar ON Musteriler.musteri_id = Araclar.musteri_id
                  WHERE Musteriler.ad LIKE ? OR Araclar.plaka LIKE ? OR Araclar.marka LIKE ?
                  ORDER BY {siralama_sql}
                  LIMIT ? OFFSET ?''', (aranan_sart, aranan_sart, aranan_sart, sayfa_basina_kayit, offset_degeri))
    else:
        c.execute(f'''
                  SELECT DISTINCT Musteriler.musteri_id
                  FROM Musteriler
                  LEFT JOIN Araclar ON Musteriler.musteri_id = Araclar.musteri_id
                  ORDER BY {siralama_sql}
                  LIMIT ? OFFSET ?''', (sayfa_basina_kayit, offset_degeri))

    sayfa_musteri_idleri = [r[0] for r in c.fetchall()]

    gruplu_musteriler = {}
    if sayfa_musteri_idleri:
        yer_tutucular = ','.join(['?'] * len(sayfa_musteri_idleri))
        c.execute(f'''
                  SELECT
                    Musteriler.ad,
                    Musteriler.soyad,
                    Musteriler.telefon,
                    Araclar.marka,
                    Araclar.plaka,
                    Musteriler.musteri_id,
                    Araclar.arac_id
                  FROM Musteriler
                  LEFT JOIN Araclar ON Musteriler.musteri_id = Araclar.musteri_id
                  WHERE Musteriler.musteri_id IN ({yer_tutucular})''', sayfa_musteri_idleri)
        birlesik_liste = c.fetchall()

        for kayit in birlesik_liste:
            ad, soyad, telefon, marka, plaka, musteri_id, arac_id = kayit

            if musteri_id not in gruplu_musteriler:
                gruplu_musteriler[musteri_id] = {
                    'musteri_id': musteri_id,
                    'ad': ad,
                    'soyad': soyad,
                    'telefon': telefon,
                    'araclar': []
                }

            if arac_id:
                tahmini_km, kalan_km, durum, dusuk_guven, sebep = arac_durumu_hesapla(c, arac_id, model)
                gruplu_musteriler[musteri_id]['araclar'].append({
                    'arac_id': arac_id,
                    'marka': marka,
                    'plaka': plaka,
                    'tahmini_km': tahmini_km,
                    'kalan_km': kalan_km,
                    'durum': durum,
                    'dusuk_guven': dusuk_guven,
                    'sebep': sebep
                })

    zenginlestirilmis_liste = [gruplu_musteriler[mid] for mid in sayfa_musteri_idleri if mid in gruplu_musteriler]
    c.execute('''
              SELECT Musteriler.ad, Musteriler.soyad, Araclar.marka, Araclar.plaka, Araclar.arac_id
              FROM Musteriler
              JOIN Araclar ON Musteriler.musteri_id = Araclar.musteri_id
              ''')
    tum_araclar = c.fetchall()

    uyari_listesi = []
    for ad, soyad, marka, plaka, arac_id in tum_araclar:
        tahmini_km, kalan_km, durum, dusuk_guven, sebep = arac_durumu_hesapla(c, arac_id, model)
        if durum in ('kritik', 'yaklasiyor'):
            uyari_listesi.append([ad, soyad, None, marka, plaka, None, arac_id, tahmini_km, kalan_km, durum, sebep])

    uyari_listesi.sort(key=lambda x: (x[9] != 'kritik', x[8]))

    c.execute('''SELECT lastik_tipi FROM Islemler''')
    tum_lastik_tipleri = [r[0].strip().upper() for r in c.fetchall()]

    sayac = Counter(tum_lastik_tipleri)
    lastik_oranlari = list(sayac.items())

    grafik_filtresi = request.args.get('filtre', 'hepsi')

    if grafik_filtresi == '6ay':
        c.execute('''
                    SELECT STRFTIME('%Y-%m', islem_tarihi), COUNT(*)
                    FROM Islemler
                    WHERE islem_tarihi >= date('now', '-6 months')
                    GROUP BY STRFTIME('%Y-%m', islem_tarihi)
                    ORDER BY STRFTIME('%Y-%m', islem_tarihi) ASC
                    ''')
    elif grafik_filtresi == '1yil':
        c.execute('''
                    SELECT STRFTIME('%Y-%m', islem_tarihi), COUNT(*)
                    FROM Islemler
                    WHERE islem_tarihi >= date('now', '-1 year')
                    GROUP BY STRFTIME('%Y-%m', islem_tarihi)
                    ORDER BY STRFTIME('%Y-%m', islem_tarihi) ASC
                    ''')
    else:
        c.execute('''
                    SELECT STRFTIME('%Y-%m', islem_tarihi), COUNT(*)
                    FROM Islemler
                    GROUP BY STRFTIME('%Y-%m', islem_tarihi)
                    ORDER BY STRFTIME('%Y-%m', islem_tarihi) ASC
                    ''')
    
    islem_tarihleri = c.fetchall()

    try:
        with open('model_bilgisi.json', 'r', encoding='utf-8') as f:
            model_bilgisi = json.load(f)
    except FileNotFoundError:
        model_bilgisi = None

    conn.close()

    return render_template('index.html', liste=zenginlestirilmis_liste,
                            toplam_musteri=toplam_musteri,
                            toplam_arac=toplam_arac,
                            toplam_islem=toplam_islem,
                            current_page=sayfa,
                            total_pages=toplam_sayfa,
                            lastik_oranlari=lastik_oranlari,
                            islem_tarihleri=islem_tarihleri,
                            grafik_filtresi=grafik_filtresi,
                            uyari_listesi=uyari_listesi,
                            model_bilgisi=model_bilgisi,
                            arama_kelimesi=gelen_arama,
                            sirala=sirala)
    
@app.route('/yeni-musteri', methods=['GET', 'POST'])
def yeni_musteri():
    if 'giris_yapildi' not in session:
        return redirect('/login')

    if request.method == 'POST':
        gelen_ad = request.form.get('isim')
        gelen_soyad = request.form.get('soyisim')
        gelen_telefon = request.form.get('telefon')
        gelen_email = request.form.get('email')

        hatalar = []
        if not gelen_ad or len(gelen_ad) < 2:
            hatalar.append("İsim en az 2 karakter olmalıdır!")
        if not gelen_soyad or len(gelen_soyad) < 2:
            hatalar.append("Soyisim en az 2 karakter olmalıdır.")
        
        telefon_temiz = re.sub(r'\s+', '', gelen_telefon)
        if not re.match(r'^0\d{10}$', telefon_temiz):
            hatalar.append("Telefon numarası 05XX XXX XX XX formatında (11 haneli) olmalı.")

        if not gelen_email or not re.match(r'^[^@\s]+@[^@\s]+\.[^@\s]+$', gelen_email):
            hatalar.append("Lütfen geçerli bir e-posta adresi giriniz.")
        
        if hatalar:
            for hata in hatalar:
                flash(hata, "danger")
            return redirect('/yeni-musteri')
        
        conn = sqlite3.connect('database.db')
        c = conn.cursor()
        c.execute('''INSERT INTO Musteriler (ad, soyad, telefon, email) VALUES (?, ?, ?, ?)''',(gelen_ad, gelen_soyad, gelen_telefon, gelen_email))
        conn.commit()
        conn.close()

        flash("Müşteri başarıyla eklendi!", "success")
        return redirect('/')
    else:
        return render_template('musteri_ekle.html')
    
@app.route('/yeni-arac', methods=['GET', 'POST'])
def yeni_arac():
    if 'giris_yapildi' not in session:
        return redirect('/login')
    
    if request.method == 'POST':
        gelen_marka = request.form.get('marka')
        gelen_model = request.form.get('model')
        gelen_plaka = request.form.get('plaka')
        gelen_musteri_id = request.form.get('musteri_id')

        hatalar = []
        if not gelen_marka or len(gelen_marka) < 2:
            hatalar.append("Marka en az 2 karakter olmalıdır.")
        if not gelen_model or len(gelen_model) <1:
            hatalar.append("Model boş bırakılamaz.")
        if not gelen_musteri_id:
            hatalar.append("Lütfen bir müşteri seçin.")
        
        plaka_temiz = gelen_plaka.replace(' ', '')
        if not re.match(r'^\d{2}[A-ZÇĞİÖŞÜ]{1,3}\d{2,4}$', plaka_temiz):
            hatalar.append("Plaka formatı geçersiz. Örnek: 34 ABC 123")

        if hatalar:
            for hata in hatalar:
                flash(hata, "danger")
            return redirect('/yeni-arac')

        conn = sqlite3.connect('database.db')
        c = conn.cursor()
        c.execute('''INSERT INTO Araclar (marka, model, plaka, musteri_id) VALUES(?, ?, ?, ?)''', (gelen_marka, gelen_model, gelen_plaka, gelen_musteri_id))
        conn.commit()
        conn.close()

        flash("Yeni araç başarıyla sisteme eklendi!", "success")
        return redirect('/')
    
    conn = sqlite3.connect('database.db')
    c = conn.cursor()
    c.execute('''SELECT * FROM Musteriler''')
    tum_musteriler = c.fetchall()
    conn.close()
    return render_template('arac_ekle.html', musteri_listesi=tum_musteriler)

@app.route('/yeni-islem', methods=['GET', 'POST'])
def yeni_islem():
    if 'giris_yapildi' not in session:
        return redirect('/login')
    
    if request.method == 'POST':
        arac_id = request.form.get('arac_id')
        islem_tarihi = request.form.get('islem_tarihi')
        km = request.form.get('kilometre')
        lastik_tipi = request.form.get('lastik_tipi')

        hatalar = []

        if islem_tarihi:
            try:
                girilen_tarih = datetime.strptime(islem_tarihi, '%Y-%m-%d')
                if girilen_tarih.date() > datetime.today().date():
                    hatalar.append("İşlem tarihi gelecekte bir tarih olamaz.")
            except ValueError:
                hatalar.append("Geçersiz tarih formatı.")
        else:
            hatalar.append("Lütfen işlem tarihi seçin.")

        try:
            km_sayi = int(km)
            if km_sayi < 0:
                flash("Hata: Kilometre eksi bir değer olamaz!", "danger")
                return redirect(request.referrer or '/') 

            conn_kontrol = sqlite3.connect('database.db')
            c_kontrol = conn_kontrol.cursor()
            c_kontrol.execute("SELECT MAX(kilometre) FROM Islemler WHERE arac_id = ?", (arac_id,))
            max_km = c_kontrol.fetchone()[0]
            conn_kontrol.close()

            if max_km is not None and km_sayi <= max_km:
                flash(f"Hata: Girilen kilometre ({km_sayi}), sistemdeki son işlem kilometresinden ({max_km}) küçük veya ona eşit olamaz!", "danger")
                return redirect(request.referrer or '/')
            
        except ValueError:
            flash("Hata: Lütfen kilometre alanına sadece sayısal bir değer giriniz!", "danger")
            return redirect(request.referrer or '/')
        
        if hatalar:
            for hata in hatalar:
                flash(hata, "danger")
            return redirect(request.referrer or '/')

        conn = sqlite3.connect('database.db')
        c = conn.cursor()
        c.execute('''INSERT INTO Islemler (arac_id, islem_tarihi, kilometre, lastik_tipi) VALUES (?, ?, ?, ?)''', (arac_id, islem_tarihi, km, lastik_tipi)) 
        conn.commit()
        conn.close()

        flash("Yeni işlem kaydı başarıyla eklendi!", "success")
        return redirect('/')

    gelen_secili_arac = request.args.get('secili_arac')

    conn = sqlite3.connect('database.db')
    c = conn.cursor()
    c.execute('''SELECT * FROM Araclar''')
    arac_listesi = c.fetchall()
    conn.commit()
    conn.close()
    return render_template('islem_ekle.html', arac_listesi=arac_listesi, secili_arac=gelen_secili_arac)
  
@app.route('/islem-sil/<silinecek_id>')
def islem_sil(silinecek_id):
    if 'giris_yapildi' not in session:
        return redirect('/login')
    
    conn = sqlite3.connect('database.db')
    c = conn.cursor()
    c.execute('''DELETE FROM Islemler WHERE islem_id = ?''',(silinecek_id,))

    conn.commit()
    conn.close()

    flash("İşlem geçmişi başarıyla silindi.", "success")
    return redirect('/')

@app.route('/musteri-sil/<silinecek_id>')
def musteri_sil(silinecek_id):
    if 'giris_yapildi' not in session:
        return redirect('/login')
    
    conn = sqlite3.connect('database.db')
    c = conn.cursor()

    c.execute("PRAGMA foreign_keys = ON;")
    c.execute('''DELETE FROM Musteriler WHERE musteri_id = ?''', (silinecek_id,))

    conn.commit()
    conn.close()

    flash("Müşteri ve ona ait tüm kayıtlar başarıyla silindi!", "success")
    return redirect('/')

@app.route('/arac-sil/<silinecek_id>')
def arac_sil(silinecek_id):
    if 'giris_yapildi' not in session:
        return redirect('/login')
    
    conn = sqlite3.connect('database.db')
    c = conn.cursor()
    
    c.execute("PRAGMA foreign_keys = ON;")
    c.execute('''DELETE FROM Araclar WHERE arac_id = ?''', (silinecek_id,))

    conn.commit()
    conn.close()

    flash("Araç sistemden başarıyla silindi.", "success")
    return redirect('/')

@app.route('/musteri-duzenle/<id>', methods=['GET', 'POST'])
def musteri_duzenle(id):
    if 'giris_yapildi' not in session:
        return redirect('/login')
    
    if request.method == 'POST':
        yeni_ad = request.form.get('isim').strip()
        yeni_soyad = request.form.get('soyisim').strip()
        yeni_telefon = request.form.get('telefon').strip()
        yeni_email = request.form.get('email', '').strip()

        hatalar = []
        if not yeni_ad or len(yeni_ad) < 2:
            hatalar.append("İsim en az 2 karakter olmalıdır!")
        if not yeni_soyad or len(yeni_soyad) < 2:
            hatalar.append("Soyisim en az 2 karakter olmalıdır.")
        
        telefon_temiz = re.sub(r'\s+', '', yeni_telefon)
        if not re.match(r'^0\d{10}$', telefon_temiz):
            hatalar.append("Telefon numarası 05XX XXX XX XX formatında (11 haneli) olmalı.")
        
        if not yeni_email or not re.match(r'^[^@\s]+@[^@\s]+\.[^@\s]+$', yeni_email):
            hatalar.append("Lütfen geçerli bir e-posta adresi giriniz.")
        
        if hatalar:
            for hata in hatalar:
                flash(hata, "danger")
            return redirect(f'/musteri-duzenle/{id}')

        conn = sqlite3.connect('database.db')
        c = conn.cursor()

        c.execute('''UPDATE Musteriler 
                  SET ad = ?, soyad = ?, telefon = ?, email = ? 
                  WHERE musteri_id = ?''',(yeni_ad, yeni_soyad, yeni_telefon, yeni_email, id))
        conn.commit()
        conn.close()

        flash("Müşteri bilgileri güncellendi!", "success")
        return redirect('/')
    else:
        conn = sqlite3.connect('database.db')
        c = conn.cursor()
        c.execute('''SELECT * FROM Musteriler WHERE musteri_id = ?''', (id,))
        secilen_musteri = c.fetchone()
        conn.commit()
        conn.close()
        return render_template('musteri_duzenle.html', musteri=secilen_musteri)
    
@app.route('/arac-duzenle/<id>',  methods = ['GET', 'POST'])
def arac_duzenle(id):
    if 'giris_yapildi' not in session:
        return redirect('/login')
    
    if request.method == 'POST':
        yeni_musteri_id = request.form.get('musteri_id')
        yeni_marka = request.form.get('marka').strip()
        yeni_model = request.form.get('model').strip()
        yeni_plaka = request.form.get('plaka').strip().upper()

        hatalar = []
        if not yeni_marka or len(yeni_marka) < 2:
            hatalar.append("Marka en az 2 karakter olmalıdır.")
        if not yeni_model or len(yeni_model) < 1:
            hatalar.append("Model boş bırakılamaz.")
        if not yeni_musteri_id:
            hatalar.append("Lütfen bir müşteri seçin.")

        plaka_temiz = yeni_plaka.replace(' ', '')
        if not re.match(r'^\d{2}[A-ZÇĞİÖŞÜ]{1,3}\d{2,4}$', plaka_temiz):
            hatalar.append("Plaka formatı geçersiz. Örnek: 34 ABC 123")

        if hatalar:
            for hata in hatalar:
                flash(hata, "danger")
            return redirect(f'/arac-duzenle/{id}')

        conn = sqlite3.connect('database.db')
        c = conn.cursor()

        c.execute('''UPDATE Araclar 
                  SET musteri_id = ?, marka = ?, model = ?, plaka = ?
                  WHERE arac_id = ?''', (yeni_musteri_id, yeni_marka, yeni_model, yeni_plaka, id))
        conn.commit()
        conn.close()

        flash("Araç bilgileri güncellendi.", "success")
        return redirect('/')
    else:
        conn = sqlite3.connect('database.db')
        c = conn.cursor()
        c.execute('''SELECT * FROM Araclar WHERE arac_id = ?''', (id,))
        secilen_arac = c.fetchone()
        
        c.execute('''SELECT * FROM Musteriler''')
        tum_musteriler = c.fetchall()

        conn.commit()
        conn.close()
        return render_template('arac_duzenle.html', arac=secilen_arac, musteri_listesi=tum_musteriler)
    
@app.route('/islem-duzenle/<id>', methods=['GET', 'POST'])
def islem_duzenle(id):
    if 'giris_yapildi' not in session:
        return redirect('/login')
    
    if request.method == 'POST':
        yeni_islem_tarihi = request.form.get('islem_tarihi')
        yeni_km = request.form.get('kilometre')
        yeni_lastik_tipi = request.form.get('lastik_tipi')
        yeni_arac_id = request.form.get('arac_id')

        hatalar = []
        if not yeni_arac_id:
            hatalar.append("Lütfen bir araç seçin.")

        if yeni_islem_tarihi:
            try:
                girilen_tarih = datetime.strptime(yeni_islem_tarihi, '%Y-%m-%d')
                if girilen_tarih.date() > datetime.today().date():
                    hatalar.append("İşlem tarihi gelecekte bir tarih olamaz.")
            except ValueError:
                hatalar.append("Geçersiz tarih formatı.")
        else:
            hatalar.append("Lütfen işlem tarihi seçin.")

        try:
            km_sayi = int(yeni_km)
            if km_sayi < 0:
                hatalar.append("Kilometre eksi bir değer olamaz.")
        except (ValueError, TypeError):
            hatalar.append("Lütfen kilometre alanına sadece sayısal bir değer giriniz.")

        if hatalar:
            for hata in hatalar:
                flash(hata, "danger")
            return redirect(f'/islem-duzenle/{id}')

        conn = sqlite3.connect('database.db')
        c = conn.cursor()
        
        c.execute('''UPDATE Islemler
                  SET arac_id = ?, islem_tarihi = ?, kilometre = ?, lastik_tipi = ?
                  WHERE islem_id = ?''', (yeni_arac_id, yeni_islem_tarihi, yeni_km, yeni_lastik_tipi, id))
        conn.commit()
        conn.close()

        flash("İşlem kaydı başarıyla güncellendi.", "success")
        return redirect('/')
    else:
        conn = sqlite3.connect('database.db')
        c = conn.cursor()
        c.execute('''SELECT * FROM Islemler WHERE islem_id = ?''', (id,))
        secilen_islem = c.fetchone()

        c.execute('''SELECT * FROM Araclar''')
        arac_listesi = c.fetchall()

        conn.commit()
        conn.close()
        
        return render_template('islem_duzenle.html', islem=secilen_islem, arac_listesi=arac_listesi)
    
@app.route('/islem-gecmisi/<id>')
def islem_gecmisi(id):
    if 'giris_yapildi' not in session:
        return redirect('/login')
    
    conn = sqlite3.connect('database.db')
    c = conn.cursor()

    c.execute('''
              SELECT Araclar.marka, Araclar.model, Araclar.plaka, Musteriler.ad, Musteriler.soyad
              FROM Araclar
              JOIN Musteriler ON Araclar.musteri_id = Musteriler.musteri_id
              WHERE Araclar.arac_id = ?
              ''', (id,))
    secilen_arac_bilgisi = c.fetchone()

    c.execute('''
              SELECT islem_id, islem_tarihi, kilometre, lastik_tipi
              FROM Islemler
              WHERE arac_id = ?
              ORDER BY islem_tarihi DESC
              ''', (id,))
    yapilan_islemler = c.fetchall()
    
    bugun = datetime.now()
    onceki_km = 0
    zengin_islemler = []

    for islem in reversed(yapilan_islemler):
        islem_id = islem[0]
        tarih = islem[1]
        suanki_km = islem[2]
        tip = islem[3]

        if onceki_km == 0:
            yapilan_km = 0
        else:
            yapilan_km = suanki_km - onceki_km
        onceki_km = suanki_km

        islem_tarihi_obj = datetime.strptime(tarih, '%Y-%m-%d')
        gecen_gun = (bugun - islem_tarihi_obj).days

        if gecen_gun < 30:
            zaman_mesaji = f"{gecen_gun} gün önce"
        else:
            ay_farki = gecen_gun // 30
            zaman_mesaji = f"{ay_farki} ay önce"
        
        zengin_islemler.append({
            'id': islem_id,
            'tarih': tarih,
            'km': suanki_km,
            'tip': tip,
            'fark_km': yapilan_km,
            'zaman_mesaji': zaman_mesaji
        })

    zengin_islemler.reverse()

    conn.close()
    return render_template('islem_gecmisi.html', arac=secilen_arac_bilgisi, islemler=zengin_islemler, arac_id=id)
    
ADMIN_KULLANICI = 'admin'
ADMIN_SIFRE_HASH = 'scrypt:32768:8:1$8LG8O69Yj8RU8ALL$7af945a49585bfa4a65c3cd9973bd44c83f48e1568c65c805306936d05370f567f8be33ff7cbb4b11c211e1138d117dedf00b796fa172d21bbb48cd7becb785d'

@app.route('/login', methods = ['GET', 'POST'])
def login():
    if request.method == 'POST':
        kullanici_adi = request.form.get('kullanici_adi')
        sifre = request.form.get('sifre')

        if kullanici_adi == ADMIN_KULLANICI and check_password_hash(ADMIN_SIFRE_HASH, sifre):
            session['giris_yapildi'] = True
            return redirect('/')
        else:
            flash("Kullanıcı adı veya şifre hatalı!", "danger")
            return redirect('/login')
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    flash("Sistemden güvenli bir şekilde çıkış yaptınız.", "success")
    return redirect('/login')

@app.route('/musteri-export')
def musteri_export():
    if 'giris_yapildi' not in session:
        return redirect('/login')
    
    conn = sqlite3.connect('database.db')
    sorgu = '''
        SELECT
            Musteriler.musteri_id AS "Müşteri ID",
            Musteriler.ad AS "Ad",
            Musteriler.soyad AS "Soyad",
            Musteriler.telefon AS "Telefon",
            Musteriler.email AS "E-posta",
            Araclar.marka AS "Marka",
            Araclar.model AS "Model",
            Araclar.plaka AS "plaka"
        FROM Musteriler
        LEFT JOIN Araclar ON Musteriler.musteri_id = Araclar.musteri_id
            '''
    
    df = pd.read_sql_query(sorgu, conn)
    conn.close()

    output = io.BytesIO() # RAM'de geçici buffer
    
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Müşteriler')

        worksheet = writer.sheets['Müşteriler']
        
        for i, kolon in enumerate(df.columns, start=1): 
            
            baslik_uzunlugu = len(str(kolon))
            
            if not df[kolon].dropna().empty:
                icerik_uzunlugu = df[kolon].dropna().astype(str).str.len().max()
            else:
                icerik_uzunlugu = 0
                
            optimum_genislik = max(baslik_uzunlugu, int(icerik_uzunlugu)) + 4
            
            sutun_harfi = get_column_letter(i)
            
            worksheet.column_dimensions[sutun_harfi].width = optimum_genislik

    output.seek(0) #imleci başa sarar

    dosya_adi = f"musteri_listesi_{datetime.today().strftime('%Y-%m-%d')}.xlsx"

    return send_file(
        output,
        as_attachment=True,
        download_name=dosya_adi,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

@app.route('/islem-export/<arac_id>')
def islem_export(arac_id):
    if 'giris_yapildi' not in session:
        return redirect('/login')
    
    conn = sqlite3.connect('database.db')
    
    arac_bilgi = conn.execute('''
                              SELECT Araclar.marka, Araclar.model, Araclar.plaka, Musteriler.ad, Musteriler.soyad
                              FROM Araclar
                              JOIN Musteriler ON Araclar.musteri_id = Musteriler.musteri_id
                              WHERE Araclar.arac_id = ?
                              ''', (arac_id,)).fetchone()
    sorgu = '''
        SELECT
            islem_tarihi AS "İşlem Tarihi",
            kilometre AS "Kilometre",
            lastik_tipi AS "Lastik Tipi"
        FROM Islemler
        WHERE arac_id = ?
        ORDER BY islem_tarihi ASC
        '''
    
    df = pd.read_sql_query(sorgu, conn, params=(arac_id,))
    conn.close()

    if df.empty:
        flash("Bu araç için işlem geçmişi bulunamadı.", "warning")
        return redirect(request.referrer or '/')

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='İşlem Geçmişi')

    output = io.BytesIO() 
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='İşlem Geçmişi')

        worksheet = writer.sheets['İşlem Geçmişi']
        
        for i, kolon in enumerate(df.columns, start=1): 
            
            baslik_uzunlugu = len(str(kolon))
            
            if not df[kolon].dropna().empty:
                icerik_uzunlugu = df[kolon].dropna().astype(str).str.len().max()
            else:
                icerik_uzunlugu = 0
                
            optimum_genislik = max(baslik_uzunlugu, int(icerik_uzunlugu)) + 4
            
            sutun_harfi = get_column_letter(i)
            
            worksheet.column_dimensions[sutun_harfi].width = optimum_genislik

    output.seek(0)

    if arac_bilgi:
        plaka = arac_bilgi[2]
        dosya_adi = f"islem_gecmisi_{plaka}_{datetime.today().strftime('%Y-%m-%d')}.xlsx"
    else:
        dosya_adi = f"islem_gecmisi_{arac_id}.xlsx"
    
    return send_file(
        output,
        as_attachment=True,
        download_name=dosya_adi,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

@app.route('/model-yeniden-egit')
def model_yeniden_egit():
    if 'giris_yapildi' not in session:
        return redirect('/login')
    
    global model

    conn = sqlite3.connect('database.db')
    df = pd.read_sql_query("SELECT arac_id, islem_tarihi, kilometre, lastik_tipi FROM Islemler", conn)
    conn.close()

    if len(df) < 10:
        flash("Modeli yeniden eğitmek için en az 10 işlem kaydı gerekiyor.", "warning")
        return redirect('/')
    
    df['lastik_tipi'] = df['lastik_tipi'].apply(lastik_tipi_kodla)
    df = df.dropna(subset=['lastik_tipi'])

    df['islem_tarihi'] = pd.to_datetime(df['islem_tarihi'])
    bugun = pd.to_datetime(datetime.today().strftime('%Y-%m-%d'))
    df['gun_farki'] = (bugun - df['islem_tarihi']).dt.days

    df = df.sort_values(by=['arac_id', 'islem_tarihi'])
    df['hedef_km'] = df.groupby('arac_id')['kilometre'].shift(-1)
    df['hedef_fark'] = df['hedef_km'] - df['kilometre']
    df = df.dropna()

    y = df['hedef_fark']
    X = df.drop(['arac_id', 'islem_tarihi', 'hedef_km', 'hedef_fark'], axis=1)

    X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.2, random_state=42)
    yeni_model = RandomForestRegressor(n_estimators=100, random_state=42)
    yeni_model.fit(X_train, y_train)

    tahminler = yeni_model.predict(X_test)
    hata_payi = mean_absolute_error(y_test, tahminler)

    joblib.dump(yeni_model, 'lastik_tahmin_modeli.joblib')
    model = yeni_model #bellekteki modeli de güncelle

    model_bilgisi_yeni = {
        "mae": round(hata_payi, 2),
        "egitim_tarihi": datetime.today().strftime('%Y-%m-%d'),
        "veri_sayisi": len(df)
    }

    with open('model_bilgisi.json', 'w', encoding='utf-8') as f:
        json.dump(model_bilgisi_yeni, f, ensure_ascii=False)
    
    flash(f"Model başarıyla yeniden eğitildi! ({len(df)} kayıt kullanıldı)", "success")
    return redirect('/')

@app.route('/uyari-gonder/<arac_id>')
def uyari_gonder(arac_id):
    if 'giris_yapildi' not in session:
        return redirect('/login')

    conn = sqlite3.connect('database.db')
    c = conn.cursor()
    c.execute('''
              SELECT Musteriler.ad, Musteriler.soyad, Araclar.plaka, Musteriler.email
              FROM Araclar
              JOIN Musteriler ON Araclar.musteri_id = Musteriler.musteri_id
              WHERE Araclar.arac_id = ?
              ''', (arac_id,))
    bilgiler = c.fetchone()

    if not bilgiler:
        flash("Araç bulunamadı!", "danger")
        return redirect('/')
    
    musteri_ad = f"{bilgiler[0]} {bilgiler[1]}"
    plaka = bilgiler[2]

    tahmini_km, kalan_km, durum, dusuk_guven, sebep = arac_durumu_hesapla(c, arac_id, model)
    conn.close()

    musteri_email = bilgiler[3] if bilgiler[3] else "kayitsiz_musteri@sistem.com"

    basarili_mi, sonuc_mesaji = musteriye_uyari_maili_gonder(musteri_ad, plaka, kalan_km, musteri_email, sebep, durum)

    if basarili_mi:
        flash(f"{plaka} plakalı araç için {sonuc_mesaji} adresine uyarı maili gönderildi!", "success")
    else:
        flash(f"Mail gönderilirken bir hata oluştu. Hata: {sonuc_mesaji}", "danger")
    
    return redirect('/')

@app.errorhandler(404)
def sayfa_bulunamadi(e):
    return render_template('hata.html', kod=404, mesaj="Aradığınız sayfa bulunamadı."), 404

@app.errorhandler(500)
def sunucu_hatasi(e):
    return render_template('hata.html', kod=500, mesaj="Sunucuda beklenmeyen bir hata oluştu."), 500

@app.route('/sms-gonder/<int:arac_id>/<durum>/<sebep>')
def manuel_sms_gonder(arac_id, durum, sebep):
    test_telefon = "+905343937378"
    if sebep == 'zaman':
        if durum == 'kritik':
            mesaj = "aracınızın lastik değişim zamanı geçmiştir. KM sınırını doldurmasanız bile sürüş güvenliğiniz için randevu alınız."
        else:
            mesaj = "periyodik lastik değişim zamanınız yaklaşmaktadır. Şimdiden planlama yapmanızı öneririz."

    else:
        if durum == 'kritik':
            mesaj = "Yapay zeka analizimize göre aracınızın lastik değişim kilometresi kritik seviyededir. Lütfen servis randevusu alınız."
        else:
            mesaj = "Sent from your Twilio trial account - Test message 1234"

    basarili_mi = sms_gonder(test_telefon, mesaj)

    if basarili_mi:
        flash(f"Müşteriye SMS başarıyla gönderildi!", "success")
    else:
        flash("SMS gönderilirken bir hata oluştu.", "danger")

    return redirect(request.referrer or '/')

        
                        
if __name__ == '__main__':
    app.run(debug=True, use_reloader=False)